import asyncio
import os
import re
import sys
import subprocess
import time
import urllib.request
import unicodedata
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenPyxlImage
from PIL import Image as PILImage

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

from bs4 import BeautifulSoup
from playwright.async_api import async_playwright

# ============================================================
# CẤU HÌNH
# ============================================================
TARGET_URL       = ["https://www.facebook.com/groups/611318877397082", "https://www.facebook.com/groups/chothuevinhomes.vn/", "https://www.facebook.com/groups/347283300103313/", "https://www.facebook.com/groups/745457997222988/"]
SCRIPT_DIR       = os.path.dirname(os.path.abspath(__file__))
OUTPUT_EXCEL     = os.path.join(SCRIPT_DIR, "danh_sach_bat_dong_san.xlsx")
IMAGE_DIR        = os.path.join(SCRIPT_DIR, "downloaded_images")
CHROME_PORT      = 9222
DAYS_LIMIT       = 30  # Lấy bài trong vòng 30 ngày (1 tháng)

CHROME_PATHS = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    os.path.expandvars(r"%LOCALAPPDATA%\Google\Chrome\Application\chrome.exe"),
]

os.makedirs(IMAGE_DIR, exist_ok=True)

def find_chrome():
    for path in CHROME_PATHS:
        if os.path.exists(path):
            return path
    return None

def clean_facebook_noise(text: str) -> str:
    """Xóa triệt để các rác hệ thống Facebook, giữ đúng nội dung văn bản gốc bài đăng."""
    if not text:
        return ""
    text = unicodedata.normalize('NFC', text)
    
    # 1. Loại bỏ các mẫu văn bản rác của giao diện Facebook
    noise_patterns = [
        r"Facebook",
        r"Đã chia sẻ (với Nhóm công khai|bài viết)?",
        r"Xem thêm",
        r"Có thể là hình ảnh về.*",
        r"Thích\s+Bình luận\s+Chia sẻ",
        r"Tất cả bình luận.*",
        r"Viết bình luận công khai.*",
        r"Ảnh từ bài viết của.*",
        r"\d+:\d+\s*/\s*\d+:\d+",                 # Thời gian video (0:00 / 0:40)
        r"[a-zA-Z0-9_-]+\.com",                    # Domain rác
        r"\b[a-zA-Z0-9]{15,}\b",                   # Chuỗi ID/Mã ngẫu nhiên
    ]
    for pattern in noise_patterns:
        text = re.sub(pattern, " ", text, flags=re.IGNORECASE)
        
    # 2. Xóa các dòng rác chứa thông tin tương tác
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    valid_lines = []
    
    for line in lines:
        if re.match(r'^(thích|bình luận|chia sẻ|gửi tin nhắn|chi tiết|chia sẻ bài viết)$', line, re.IGNORECASE):
            continue
        valid_lines.append(line)
        
    cleaned = " ".join(valid_lines)
    cleaned = re.sub(r'\s+', ' ', cleaned).strip()

    # 3. Lọc bỏ tên người đăng ở đầu
    cleaned = re.sub(
        r'^[A-ZĐÀÁÂÃÈÉÊÌÍÒÓÔÕÙÚĂĐĨŨƠàáâãèéêìíòóôõùúăđĩũơƯĂẠẢẤẦẨẪẬẮẰẲẴẶẸẺẼỀỀỂưăạảấầnẩẫậnắằẳẵặẹẻẽềềểỄỆỈỊỌỎỐỒỔỖỘỚỜỞỠỢỤỦỨỪễệỉịọỏốồ ổỗộớờởỡợụủứừỬỮỰỲỴÝỶỸửữựỳỵỷỹ\s]{2,25}\s*[-·•]?\s*(?=(cho thuê|pass|căn|bán|chính chủ|nhà|phòng|quỹ|giỏ hàng|suất|trực tiếp|danh sách))', 
        '', 
        cleaned, 
        flags=re.IGNORECASE
    )

    return cleaned.strip()

def create_property_fingerprint(cleaned_desc: str) -> str:
    """Tạo dấu vân tay độc nhất cho từng CĂN HỘ để lọc trùng giữa các người đăng khác nhau."""
    loc = re.search(r"(BE\d+|S\d+\.\d+|S\d+|ZR\d+|R\d+\.\d+|BE4|BE3|BE2|BE1|ORIGAMI|RAINBOW|MASTERI|GRAND PARK|ZENPARK|PAVILION|ZURICH|IMPERIA|OCP\d*|tòa\s+[A-Za-z0-9]+|toà\s+[A-Za-z0-9]+)", cleaned_desc, re.IGNORECASE)
    pn = re.search(r"(\d+)\s*(?:n|pn|phòng ngủ)", cleaned_desc, re.IGNORECASE)
    wc = re.search(r"(\d+)\s*(?:vs|wc|phòng tắm|vệ sinh)", cleaned_desc, re.IGNORECASE)
    price = re.search(r"(\d+(?:[\.,]\d+)?)\s*(?:triệu|trđ|tr|tỷ|ty)", cleaned_desc, re.IGNORECASE)
    
    loc_str = loc.group(0).upper().replace(" ", "") if loc else ""
    pn_str = pn.group(1) if pn else ("studio" if "studio" in cleaned_desc.lower() else "")
    wc_str = wc.group(1) if wc else ""
    price_str = price.group(0).lower().replace(" ", "") if price else ""
    
    if loc_str and (pn_str or price_str):
        return f"{loc_str}_{pn_str}_{wc_str}_{price_str}"
    
    clean_str = re.sub(r'[^\w]', '', cleaned_desc.lower())
    return clean_str[:70]

def parse_price(text: str):
    """
    Trích xuất giá chuẩn theo VNĐ.
    Trả về "" nếu bài đăng không có số giá rõ ràng (mập mờ về giá).
    """
    t = text.lower()
    
    # Bỏ qua ngay nếu có từ ngữ mập mờ về giá
    if re.search(r"(inbox|ib|thương lượng|thoả thuận|thỏa thuận|ép giá|giá mầm non|giá tiểu học|giá hạt dẻ|liên hệ|lh)", t) and not re.search(r'(\d+(?:[\.,]\d+)?)\s*(?:triệu|trđ|tr|tỷ|ty|k)', t):
        return ""

    match_full_vnd = re.search(r'(\d{1,3}(?:\.\d{3}){2,3})', t)
    if match_full_vnd:
        return int(match_full_vnd.group(1).replace('.', ''))

    match_tr = re.search(r'(\d+(?:[\.,]\d+)?)\s*(?:triệu|trđ|tr)', t)
    if match_tr:
        val_tr = float(match_tr.group(1).replace(',', '.'))
        return int(val_tr * 1_000_000)

    match_ty = re.search(r'(\d+(?:[\.,]\d+)?)\s*(?:tỷ|ty)', t)
    if match_ty:
        val_ty = float(match_ty.group(1).replace(',', '.'))
        return int(val_ty * 1_000_000_000)

    return ""

def is_within_one_month(time_str: str) -> bool:
    """Kiểm tra bài viết có thuộc 30 ngày gần đây hay không."""
    if not time_str:
        return True
    
    t = time_str.lower()
    if any(x in t for x in ["vừa xong", "phút", "giờ", "hôm nay", "hôm qua"]):
        return True
    
    match_days = re.search(r'(\d+)\s*ngày', t)
    if match_days:
        return int(match_days.group(1)) <= DAYS_LIMIT

    match_weeks = re.search(r'(\d+)\s*tuần', t)
    if match_weeks:
        return int(match_weeks.group(1)) * 7 <= DAYS_LIMIT

    match_months = re.search(r'(\d+)\s*tháng', t)
    if match_months:
        return int(match_months.group(1)) < 1

    return True

def parse_rental_post(cleaned_desc: str, image_urls: list, post_index: int, post_url: str, price_vnd: int) -> dict:
    loc_match = re.search(
        r"(BE\d+|S\d+\.\d+|S\d+|ZR\d+|R\d+\.\d+|BE4|BE3|BE2|BE1|ORIGAMI|RAINBOW|OCEAN PARK\s*\d?|"
        r"MASTERI|GRAND PARK|ZENPARK|PAVILION|ZURICH|IMPERIA|OCP\d*|tòa\s+[A-Za-z0-9]+|toà\s+[A-Za-z0-9]+)",
        cleaned_desc, re.IGNORECASE
    )
    location = loc_match.group(0).strip().upper() if loc_match else ""

    pn_match = re.search(r"(\d+)\s*(?:n|pn|phòng ngủ)", cleaned_desc, re.IGNORECASE)
    if pn_match:
        bedrooms = int(pn_match.group(1))
    elif "studio" in cleaned_desc.lower():
        bedrooms = "Studio"
    else:
        bedrooms = ""

    wc_match = re.search(r"(\d+)\s*(?:vs|wc|phòng tắm|vệ sinh)", cleaned_desc, re.IGNORECASE)
    bathrooms = int(wc_match.group(1)) if wc_match else ""

    area_match = re.search(r"(\d+(?:[\.,]\d+)?)\s*(?:m2|m²)", cleaned_desc, re.IGNORECASE)
    area = float(area_match.group(1).replace(',', '.')) if area_match else ""

    if re.search(r"(full|đầy đủ)", cleaned_desc, re.IGNORECASE):
        furniture = "Full đồ"
    elif re.search(r"(cơ bản|đồ cơ bản)", cleaned_desc, re.IGNORECASE):
        furniture = "Cơ bản"
    else:
        furniture = ""

    # TẢI TOÀN BỘ ẢNH CỦA CĂN HỘ VÀO THƯ MỤC RIÊNG
    post_img_dir = os.path.join(IMAGE_DIR, f"can_{post_index:03d}")
    os.makedirs(post_img_dir, exist_ok=True)
    
    downloaded_image_paths = []
    for idx, img_url in enumerate(image_urls, 1):
        try:
            save_path = os.path.join(post_img_dir, f"anh_{idx}.jpg")
            req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
            with urllib.request.urlopen(req, timeout=10) as response, open(save_path, 'wb') as out_file:
                out_file.write(response.read())
            downloaded_image_paths.append(save_path)
        except Exception:
            pass

    first_image_path = downloaded_image_paths[0] if downloaded_image_paths else None

    return {
        "source": "Facebook Group",
        "location": location,
        "price_vnd": price_vnd,
        "area": area,
        "furniture": furniture,
        "bedrooms": bedrooms,
        "bathrooms": bathrooms,
        "description": cleaned_desc,
        "link": post_url if post_url else TARGET_URL,
        "image_path": first_image_path,
        "total_images": len(downloaded_image_paths),
        "folder_path": post_img_dir
    }

def export_to_styled_excel(data_list, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dữ liệu"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Hình ảnh", "Số lượng ảnh", "Nguồn", "Phân khu / Dự án", "Giá (VNĐ)", 
        "Diện tích (m²)", "Nội thất", "Số phòng ngủ", 
        "Số phòng tắm", "Mô tả ngắn", "Link gốc", "Thư mục ảnh"
    ]
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for row_idx, item in enumerate(data_list, start=2):
        row_data = [
            "",
            f"{item['total_images']} ảnh",
            item["source"],
            item["location"],
            item["price_vnd"],
            item["area"],
            item["furniture"],
            item["bedrooms"],
            item["bathrooms"],
            item["description"],
            item["link"],
            item["folder_path"]
        ]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 65

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            
            if col_idx in [2, 3, 4, 7, 8, 9]:
                cell.alignment = center_align
            elif col_idx in [5, 6]:
                cell.alignment = right_align
                if isinstance(cell.value, (int, float)):
                    if col_idx == 5: cell.number_format = '#,##0'
                    elif col_idx == 6: cell.number_format = '#,##0.0'
            else:
                cell.alignment = left_align

        # CHÈN ẢNH VÀO CỘT A
        img_path = item.get("image_path")
        if img_path and os.path.exists(img_path):
            try:
                img_pil = PILImage.open(img_path)
                img_pil.thumbnail((80, 80))
                
                temp_img_path = os.path.join(IMAGE_DIR, f"thumb_{row_idx}.jpg")
                img_pil.save(temp_img_path)

                img = OpenPyxlImage(temp_img_path)
                ws.add_image(img, f"A{row_idx}")
            except Exception:
                pass

    column_widths = {
        'A': 14, 'B': 14, 'C': 16, 'D': 20, 'E': 18, 
        'F': 14, 'G': 15, 'H': 14, 'I': 14, 'J': 50, 'K': 30, 'L': 35
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(output_path)

async def main():
    chrome_path = find_chrome()
    if not chrome_path:
        print("❌ KHÔNG TÌM THẤY CHROME!")
        return

    profile_dir = os.path.join(SCRIPT_DIR, "chrome_debug_profile")
    os.makedirs(profile_dir, exist_ok=True)

    subprocess.Popen([
        chrome_path,
        f"--remote-debugging-port={CHROME_PORT}",
        f"--user-data-dir={profile_dir}",
        "--no-first-run",
        "--no-default-browser-check",
        "about:blank"
    ])

    print("Chờ Chrome khởi động (5 giây)...")
    await asyncio.sleep(5)

    async with async_playwright() as p:
        try:
            browser = await p.chromium.connect_over_cdp(f"http://localhost:{CHROME_PORT}")
        except Exception as e:
            print(f"Lỗi kết nối Chrome: {e}")
            return

        context = browser.contexts[0]
        page = await context.new_page()

        print("Đang mở Facebook Group...")
        await page.goto(TARGET_URL, wait_until="domcontentloaded", timeout=60000)
        await asyncio.sleep(3)

        if "login" in page.url:
            print("\nVUI LÒNG ĐĂNG NHẬP FACEBOOK TRÊN CỬA SỔ CHROME VỪA MỞ")
            input("[Bấm ENTER sau khi đăng nhập thành công] > ")
            await page.goto(TARGET_URL, wait_until="domcontentloaded")

        print("Đang cuộn trang liên tục (Chỉ lấy bài CÓ ẢNH & CÓ GIÁ RÕ RÀNG)...")
        
        rental_posts = []
        seen_property_fingerprints = set()
        seen_post_links = set()
        
        last_height = 0
        no_change_count = 0
        reached_time_limit = False

        while not reached_time_limit:
            await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
            await asyncio.sleep(3.5)

            html_content = await page.content()
            soup = BeautifulSoup(html_content, "html.parser")
            
            feed_blocks = soup.find_all("div", attrs={"role": "feed"})
            if feed_blocks:
                articles = feed_blocks[0].find_all("div", attrs={"data-scrolling-node": None})
            else:
                articles = soup.find_all("div", attrs={"role": "article"})

            for block in articles:
                # 1. KIỂM TRA THỜI GIAN BÀI ĐĂNG (TRONG 1 THÁNG)
                time_tag = block.find("a", attrs={"aria-label": True}) or block.find("span", text=re.compile(r'ngày|tháng|tuần|giờ|phút'))
                time_str = time_tag.get_text() if time_tag else ""
                
                if not is_within_one_month(time_str):
                    reached_time_limit = True
                    print(f"\n⏹️ Đã chạm mốc bài viết cũ hơn 30 ngày. Dừng cào dữ liệu.")
                    break

                # 2. ĐIỀU KIỆN LỌC 1: BẮT BUỘC BÀI VIẾT PHẢI CÓ ẢNH
                img_tags = block.find_all("img")
                image_urls = []
                for img in img_tags:
                    src = img.get("src")
                    if src and "scontent" in src and not re.search(r'p\d+x\d+|s32x32|s50x50|emoji', src):
                        if src not in image_urls:
                            image_urls.append(src)

                if not image_urls:
                    # Bỏ qua nếu không có ảnh nào
                    continue

                # 3. LỌC TRÙNG LINK BÀI VIẾT
                post_link = ""
                link_tag = block.find("a", href=re.compile(r'/groups/.+/posts/'))
                if link_tag:
                    post_link = link_tag.get("href", "")
                    if post_link.startswith("/"):
                        post_link = "https://www.facebook.com" + post_link
                    if post_link in seen_post_links:
                        continue
                    seen_post_links.add(post_link)

                # 4. BÓC TÁCH & LÀM SẠCH VĂN BẢN
                msg_div = block.find("div", attrs={"data-ad-preview": "message"}) or block.find("div", attrs={"dir": "auto"})
                raw_text = msg_div.get_text(separator="\n").strip() if msg_div else block.get_text(separator="\n").strip()

                cleaned_text = clean_facebook_noise(raw_text)
                if len(cleaned_text) < 20:
                    continue

                if re.search(r"(đã thuê|đã cho thuê|xoá bài|xóa bài|đã bán)", cleaned_text, re.IGNORECASE):
                    continue

                # 5. ĐIỀU KIỆN LỌC 2: LOẠI BỎ GIÁ MẬP MỜ (BẮT BUỘC PHẢI CÓ GIÁ SỐ VNĐ CỤ THỂ)
                price_vnd = parse_price(cleaned_text)
                if not price_vnd or price_vnd == "":
                    # Bỏ qua nếu không có giá hoặc giá mập mờ (inbox, thỏa thuận,...)
                    continue

                # 6. KHỬ TRÙNG TỰ ĐỘNG GIỮA CÁC NGƯỜI ĐĂNG
                fingerprint = create_property_fingerprint(cleaned_text)
                if fingerprint in seen_property_fingerprints:
                    continue
                seen_property_fingerprints.add(fingerprint)

                parsed = parse_rental_post(cleaned_text, image_urls, len(rental_posts) + 1, post_link, price_vnd)
                rental_posts.append(parsed)
                print(f"✔️ Lấy thành công căn [{len(rental_posts)}] (Giá: {price_vnd:,} VNĐ - {len(image_urls)} ảnh): {cleaned_text[:35]}...")

            new_height = await page.evaluate("document.body.scrollHeight")
            if new_height == last_height:
                no_change_count += 1
                if no_change_count >= 4:
                    print("\nĐã cuộn hết danh sách bài viết trong Group.")
                    break
            else:
                no_change_count = 0
            last_height = new_height

        await browser.close()

    if not rental_posts:
        print("⚠️ Không tìm thấy bài viết hợp lệ nào!")
        return

    # Xuất Excel
    final_output_path = OUTPUT_EXCEL
    try:
        export_to_styled_excel(rental_posts, final_output_path)
    except PermissionError:
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        final_output_path = os.path.join(SCRIPT_DIR, f"danh_sach_bat_dong_san_{timestamp}.xlsx")
        export_to_styled_excel(rental_posts, final_output_path)
        print(f"\n⚠️ CẢNH BÁO: File Excel gốc đang mở! Dữ liệu đã lưu sang file mới:")

    print("\n" + "="*60)
    print(f"✅ THÀNH CÔNG! Đã thu thập {len(rental_posts)} căn chuẩn (có ảnh + có giá rõ ràng + không trùng).")
    print(f"📁 File Excel: {final_output_path}")
    print(f"🖼️ Thư mục chứa toàn bộ ảnh: {IMAGE_DIR}")
    print("="*60)

if __name__ == "__main__":
    asyncio.run(main())